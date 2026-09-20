import csv,json,pathlib,hashlib,collections
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
P=pathlib.Path('output/eploy-sheet6-corrected-2026-09-18')
original=pathlib.Path(r'D:\Programs\Java\downloads\eploy-sheet6-1-uk-jobs.csv')
def read(p):
 with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
old=read(original);new=read(P/'eploy-sheet6-1-uk-jobs-corrected.csv');audit=read(P/'field-correction-audit.csv');summary=json.loads((P/'summary.json').read_text())
assert len(old)==len(new)==len(audit)==1528
assert list(old[0])==list(new[0])
for i,(before,after) in enumerate(zip(old,new)):
 for k in before:
  if k not in ['employmentType','worktype']:assert before[k]==after[k],(i,k)
 assert after['worktype'] in ['', 'Hybrid','Remote','On-site']
 assert not any(x in after['employmentType'] for x in ['Per Week','Not Specified','Contract:','Headhunt'])
 assert after['employmentType']==audit[i]['employmentType'] and after['worktype']==audit[i]['worktype']
# Regression checks for misleading matches and corrupted original fields.
assert new[0]['worktype']=='Hybrid'
assert new[1]['worktype']=='Remote'
assert new[261]['worktype']=='On-site' # role-specific statement beats generic benefits
assert new[505]['employmentType']=='Permanent' # not its hours-per-week field
assert new[543]['worktype']=='On-site' # office + construction site, not home-working
assert new[734]['worktype']=='' # hybrid job responsibilities, not workplace
assert new[1502]['worktype']=='' # Agile/Waterfall hybrid delivery is not a workplace
assert new[778]['employmentType']=='' # 40 hours alone is not a contract type
wb=Workbook();jobs=wb.active;jobs.title='Corrected jobs';ev=wb.create_sheet('Correction audit');info=wb.create_sheet('Read me',0)
notes=[('Eploy job-field corrections','September 18, 2026'),('Scope','Only employmentType and worktype changed. Original row order, duplicates and all 12 other columns retained.'),('Rows',1528),('Unique vacancy pages checked',1480),('Employment entries changed',summary['employmentChanged']),('Worktype values filled',sum(v for k,v in summary['worktypes'].items() if k)),('Employment type unresolved',summary['employmentTypes'].get('',0)),('Worktype unresolved',summary['worktypes'].get('',0)),('Blank values','Not established from the available evidence. A blank is not an assertion of On-site, Full-time, or Permanent.'),('Employment type','Primary vacancy contract type preferred; then source structured data; valid original labels retained where necessary. Explicit title/description clauses used only when type is otherwise missing. Numeric weekly hours are not contract types.'),('Worktype','Remote, Hybrid or On-site only when supported by workplace wording/metadata or explicit physical workplace duties. No default based on title, city or employer.'),('Conditional arrangements','Some Hybrid roles start after probation/training; conditions are retained in the audit evidence. Merely possible arrangements and generic where-possible boilerplate are not used.'),('Source limitations','One unique page returned HTTP 410. Title-mismatched pages are excluded as live evidence. Retained input values and unavailable evidence are identified in Correction audit.'),('Other fields','No correction of employer names, locations, dates, salaries or job content was requested or attempted.'),('Audit','Filter notes to find unresolved entries. row is the original CSV/Corrected jobs worksheet row number, including the header.'),('Original CSV SHA-256',hashlib.sha256(original.read_bytes()).hexdigest())]
for row in notes:info.append(row)
for ws,data in [(jobs,new),(ev,audit)]:
 ws.append(list(data[0]))
 for d in data:ws.append(list(d.values()))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='17354A');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[1].height=30
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   if isinstance(cell.value,str):cell.data_type='s'
   cell.alignment=Alignment(vertical='top',wrap_text=True)
  ws.row_dimensions[row[0].row].height=32
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=22
for col,width in {'A':14,'B':48,'C':80,'D':45,'G':35,'I':28,'J':18}.items():jobs.column_dimensions[col].width=width
for col in ['H','K','M']:ev.column_dimensions[col].width=65
for i,a in enumerate(audit,2):
 for col,field,source,evidence in [(9,'employmentType','employmentSource','employmentEvidence'),(10,'worktype','workSource','workEvidence')]:
  cell=jobs.cell(i,col)
  if not a[field]:cell.fill=PatternFill('solid',fgColor='FFF0CE')
  elif old[i-2][field]!=new[i-2][field]:cell.fill=PatternFill('solid',fgColor='E1F0E8')
  cell.comment=Comment((a[source]+': '+a[evidence]) if a[field] else 'Unresolved: source evidence does not establish this field. See Correction audit.','Source audit')
info.column_dimensions['A'].width=32;info.column_dimensions['B'].width=110
for row in info:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 row[0].font=Font(bold=True,color='17354A');info.row_dimensions[row[0].row].height=44
info.row_dimensions[1].height=30
for c in info[1]:c.fill=PatternFill('solid',fgColor='17354A');c.font=Font(bold=True,color='FFFFFF',size=14)
path=P/'eploy-sheet6-1-uk-jobs-corrected.xlsx';wb.save(path)
check=load_workbook(path,read_only=True,data_only=False)
values=list(check['Corrected jobs'].values)
assert list(values[0])==list(new[0])
assert len(values)==1529
for got,want in zip(values[1:],new):assert [(v if v is not None else '') for v in got]==list(want.values())
check.close()
result={'passed':True,'rows':1528,'columns':14,'unalteredOtherColumns':12,'csvAndWorkbookExactMatch':True,'originalSha256':hashlib.sha256(original.read_bytes()).hexdigest(),'originalMatchesCapturedInput':old==json.loads((P/'input.json').read_text(encoding='utf8'))}
assert result['originalMatchesCapturedInput']
(P/'validation.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result));print(json.dumps(summary));print(path.resolve())
