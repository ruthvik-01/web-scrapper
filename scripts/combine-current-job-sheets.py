import csv,json,hashlib,collections
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path.cwd();OUT=ROOT/'output/combined-all-jobs-2026-09-18';OUT.mkdir(exist_ok=True)
paths=[ROOT/'output/eploy-sheet6-salary-corrected-2026-09-18/eploy-sheet6-1-uk-jobs-salary-corrected.csv',ROOT/'output/fourcompany-annual-salary-corrected-2026-09-18/FourCompany-jobs-annual-salary-corrected.csv',Path(r'D:\Programs\Java\downloads\companies (4).csv')]
headers=[];combined=[];sources=[];originals=[]
for p in paths:
 digest=hashlib.sha256(p.read_bytes()).hexdigest()
 with p.open(encoding='utf-8-sig',newline='') as f:
  reader=csv.DictReader(f);cols=reader.fieldnames;rows=list(reader)
 assert cols and all(None not in row for row in rows)
 for col in cols:
  if col not in headers:headers.append(col)
 start=len(combined)+2;combined.extend(rows);originals.append((p,digest,cols,rows))
 sources.append({'path':str(p),'rows':len(rows),'firstWorksheetRow':start,'lastWorksheetRow':len(combined)+1,'sha256':digest})
normalized=[{col:r.get(col,'') for col in headers} for r in combined]
assert len(normalized)==1776 and len(headers)==15
csvpath=OUT/'all-jobs-combined.csv'
with csvpath.open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=headers);writer.writeheader();writer.writerows(normalized)
wb=Workbook();ws=wb.active;ws.title='All jobs';ws.append(headers)
for r in normalized:
 assert max(map(len,r.values()))<=32767,'Excel cell limit would truncate a source field'
 ws.append(list(r.values()))
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17354A');c.alignment=Alignment(wrap_text=True)
ws.row_dimensions[1].height=28
for row in ws.iter_rows(min_row=2):
 ws.row_dimensions[row[0].row].height=30
 for c in row:
  c.data_type='s';c.alignment=Alignment(vertical='top',wrap_text=True)
  if c.row%2==0:c.fill=PatternFill('solid',fgColor='F2F6F8')
for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
for col,width in [('A',15),('B',52),('C',85),('D',48),('G',34),('H',26),('I',26),('J',16),('K',38),('O',16)]:ws.column_dimensions[col].width=width
xlsx=OUT/'all-jobs-combined.xlsx';wb.save(xlsx)
# Verify both formats exactly, with only absent-column padding added.
with csvpath.open(encoding='utf-8-sig',newline='') as f:
 reader=csv.DictReader(f);got=list(reader);assert reader.fieldnames==headers;assert got==normalized
check=load_workbook(xlsx,read_only=True,data_only=False);assert check.sheetnames==['All jobs'];vals=list(check['All jobs'].values);assert list(vals[0])==headers;assert len(vals)==1777
for got,want in zip(vals[1:],normalized):assert [x if x is not None else '' for x in got]==list(want.values())
check.close()
offset=0
for p,digest,cols,rows in originals:
 assert hashlib.sha256(p.read_bytes()).hexdigest()==digest
 for before,after in zip(rows,normalized[offset:offset+len(rows)]):
  assert all(before[c]==after[c] for c in cols)
  assert all(after[c]=='' for c in headers if c not in cols)
 offset+=len(rows)
url_counts=collections.Counter(r['jobUrl'] for r in normalized if r['jobUrl'])
validation={'rows':1776,'columns':15,'worksheets':1,'sourceFiles':sources,'allSourceValuesAndOrderPreserved':True,'originalFilesUnchanged':True,'csvWorkbookExactMatch':True,'missingATSLeftBlank':1713,'rowsWithoutTitleRetained':sum(not r['title'] for r in normalized),'duplicateURLExtraRowsRetained':sum(n-1 for n in url_counts.values() if n>1),'exactDuplicateExtraRowsRetained':sum(n-1 for n in collections.Counter(tuple(r.values()) for r in normalized).values() if n>1)}
(OUT/'merge-validation.json').write_text(json.dumps(validation,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in validation.items() if k!='sourceFiles'},indent=2))
