import csv,re,json,pathlib,collections
from decimal import Decimal
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
P=pathlib.Path('output/eploy-sheet6-corrected-2026-09-18');D=pathlib.Path('output/eploy-sheet6-salary-corrected-2026-09-18');D.mkdir(exist_ok=True)
BASE='eploy-sheet6-1-uk-jobs-corrected';NAME='eploy-sheet6-1-uk-jobs-salary-corrected'
with (P/(BASE+'.csv')).open(encoding='utf-8-sig',newline='') as f:rows=list(csv.DictReader(f))
num=r'(\d+(?:,\d{3})*(?:\.\d+)?)(\s*[kK])?'
range_re=re.compile(r'([£$€])\s*'+num+r'\s*(?:-|–|—|\bto\b|\bup to\b)\s*([£$€])?\s*'+num)
single_re=re.compile(r'([£$€])\s*'+num)
def fmt(v):return format(v.normalize(),'f')
def amount(s,k):return Decimal(s.replace(',',''))*(1000 if k else 1)
def normalize(raw):
 s=re.sub(r'\s+',' ',raw).strip();notes=[]
 if not s:return '', 'Missing source salary'
 if re.search(r'\bNational:.*\bLondon:',s,re.I):return '', 'Multiple location-specific salary bands; no single band selected'
 m=range_re.search(s)
 if m:
  cur,lo,lk,cur2,hi,hk=m.groups()
  if cur2 and cur2!=cur:return '', 'Conflicting currencies'
  a,b=amount(lo,lk),amount(hi,hk)
  if not lk and hk:a*=1000
  elif not lk and not hk and a<1000 and b>=10000 and ',' in hi:a*=1000
  if a>b:return '', 'Inverted or ambiguous source range'
  value=f'{cur}{fmt(a)}-{cur}{fmt(b)}';notes.append('Explicit range; trailing benefits/bonus/FTE alternatives excluded')
 else:
  m=single_re.search(s)
  if m:
   cur,v,k=m.groups();value=cur+fmt(amount(v,k));notes.append('Single source amount; no second endpoint invented')
  else:
   t=re.sub(r'^(?:Details:)+','',s,flags=re.I).strip()
   m=re.fullmatch(num+r'\s*(?:\(?pro rata\)?|PA|p\.?a\.?|ph)?\s*',t,re.I)
   if not m:return '', 'No usable numeric salary (non-numeric or malformed source)'
   value='£'+fmt(amount(*m.groups()));notes.append('Numeric-only UK source; GBP assumed')
 if re.search(r'\b(?:up to|up tp)\b',s[:m.start() if hasattr(m,'start') else 0],re.I) or re.match(r'.*?(?:up to|up tp)\s*[£$€]',s,re.I):notes.append('Source upper-bound qualifier retained in audit, not an exact fixed salary')
 if re.search(r'(?:starting\s+)?from\s*[£$€]',s,re.I):notes.append('Source lower-bound qualifier retained in audit')
 if re.search(r'per hour|hourly|p/h|\d\s*ph\b',s,re.I):notes.append('Hourly amount; not annualized')
 elif re.search(r'per day|daily',s,re.I):notes.append('Daily amount; not annualized')
 if re.search(r'pro.?rata|FTE|equivalent',s,re.I):notes.append('Original pro-rata/FTE basis preserved in audit; no recalculation')
 if 'OTE' in s or re.search('on target earnings',s,re.I):notes.append('Bonus/OTE qualifiers preserved in original text; no bonus added to base pay')
 return value,'; '.join(notes)
# Focused regression checks.
for raw,expected in [
 ('£42,500-£45,000+excellent benefits','£42500-£45000'),
 ('£40,720 - £44,120 + £3,000 signing on bonus','£40720-£44120'),
 ('Base salary of £50-55k (DOE) plus on target bonus of £20k','£50000-£55000'),
 ('Base salary of £35-40,000 (DOE) plus on target bonus of £25,000','£35000-£40000'),
 ('£38-£42,000 + Bonus + Benefits','£38000-£42000'),
 ('£22,496 to £23,163 actual (£27875 to £28701 FTE)','£22496-£23163'),
 ('£13.12 - £14.12','£13.12-£14.12'),
 ('Up to £55,000 + £4,236 car allowance','£55000'),
 ('32000','£32000'),('[object Object]',''),
 ('National: £27,000 - £29,000, London: £29,500 - £32,011',''),
 ('Multi Drop Delivery Driver, Tuesday-Saturday 5am Stars','')]:
 assert normalize(raw)[0]==expected,(raw,normalize(raw),expected)
audit=[];new=[]
for i,r in enumerate(rows,2):
 value,note=normalize(r['salaryRange']);n=dict(r);n['salaryRange']=value;new.append(n)
 audit.append({'row':i,'company':r['company'],'jobId':r['jobId'],'title':r['title'],'originalSalary':r['salaryRange'],'salaryRange':value,'notes':note})
with (D/(NAME+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(new)
with (D/'salary-correction-audit.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(audit[0]));w.writeheader();w.writerows(audit)
wb=load_workbook(P/(BASE+'.xlsx'));ws=wb['Corrected jobs']
for i,(r,a) in enumerate(zip(new,audit),2):
 c=ws.cell(i,8,r['salaryRange'] or None)
 if not r['salaryRange']:c.value=None
 c.comment=Comment('Original: '+a['originalSalary']+'\n'+a['notes'],'Salary audit');c.fill=PatternFill('solid',fgColor='E1F0E8' if r['salaryRange'] else 'FFF0CE')
sheet=wb.create_sheet('Salary audit');sheet.append(list(audit[0]))
for a in audit:sheet.append(list(a.values()))
sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
for c in sheet[1]:c.fill=PatternFill('solid',fgColor='17354A');c.font=Font(bold=True,color='FFFFFF')
for col,width in [('A',10),('B',35),('C',14),('D',48),('E',80),('F',25),('G',100)]:sheet.column_dimensions[col].width=width
for row in sheet:
 sheet.row_dimensions[row[0].row].height=36
 for c in row:
  c.alignment=Alignment(wrap_text=True,vertical='top')
  if isinstance(c.value,str):c.data_type='s'
info=wb['Read me'];info['B1']='September 18, 2026 — salary formatting revision';info['B2']='This revision changes salaryRange only, preserving the prior employmentType/worktype corrections and all other values. All 1,528 rows retained.'
for row in info:
 if row[0].value=='Other fields':row[1].value='No changes to employer names, locations, dates or descriptions. Original salary wording retained in Salary audit; no currency conversion or annualization.'
info.append(['Salary format','Ranges use £30000-£40000; single values remain £30000. Source currency retained; GBP assumed only for numeric-only UK salary labels. Commas, labels and benefits removed.'])
info.append(['Salary caveats','Hourly/daily amounts retain their original numeric basis, not annualized. Up-to/from, OTE and FTE qualifiers remain in Salary audit. Non-numeric/malformed salaries and ambiguous regional bands are blank.'])
for row in list(info)[-2:]:
 info.row_dimensions[row[0].row].height=55
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(D/(NAME+'.xlsx'))
# Validate CSV roundtrip, workbook equality, and every other field including employment/worktype.
with (D/(NAME+'.csv')).open(encoding='utf-8-sig',newline='') as f:actual=list(csv.DictReader(f))
assert actual==new and len(new)==1528
pat=re.compile(r'^[£$€]\d+(?:\.\d+)?(?:-[£$€]\d+(?:\.\d+)?)?$')
for before,after in zip(rows,new):
 assert all(before[k]==after[k] for k in before if k!='salaryRange')
 assert not after['salaryRange'] or pat.fullmatch(after['salaryRange'])
check=load_workbook(D/(NAME+'.xlsx'),read_only=True);values=list(check['Corrected jobs'].values)
assert list(values[0])==list(new[0]);assert len(values)==len(new)+1
for got,want in zip(values[1:],new):assert [v if v is not None else '' for v in got]==list(want.values())
check.close()
summary={'rows':len(new),'changedSalaryCells':sum(a['salaryRange']!=b['salaryRange'] for a,b in zip(rows,new)),'ranges':sum('-' in r['salaryRange'] for r in new),'singleAmounts':sum(bool(r['salaryRange']) and '-' not in r['salaryRange'] for r in new),'blank':sum(not r['salaryRange'] for r in new),'other13ColumnsUnchanged':True,'csvWorkbookEquality':True,'parserRegressionCases':12}
(D/'salary-validation.json').write_text(json.dumps(summary,indent=2));print(json.dumps(summary))
print('Ambiguous/non-numeric input counts:',collections.Counter(a['originalSalary'] for a in audit if not a['salaryRange'] and a['originalSalary']).most_common(10))
