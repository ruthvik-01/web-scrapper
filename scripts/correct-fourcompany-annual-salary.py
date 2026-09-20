import csv,re,json,hashlib
from pathlib import Path
from decimal import Decimal
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
SRC=Path(r'D:\Programs\Java\downloads\FourCompany-jobs.csv')
OUT=Path('output/fourcompany-annual-salary-corrected-2026-09-18');OUT.mkdir(exist_ok=True)
original_hash=hashlib.sha256(SRC.read_bytes()).hexdigest()
with SRC.open(encoding='utf-8-sig',newline='') as f:old=list(csv.DictReader(f))
# Only explicitly labelled annual salary headers. Do not extract budgets, benefits,
# hourly wages, student stipends or alternative FTE figures from body paragraphs.
number=r'\d+(?:,\d{3})*(?:\.\d+)?'
pattern=re.compile(r'^(?P<indent>[^\S\r\n]*)Salary:[^\S\r\n]*(?P<bound>up to\s+|from\s+|starting from\s+)?£\s*(?P<low>'+number+r')(?:\s*(?:[-–—]|to)\s*£?\s*(?P<high>'+number+r'))?\s+(?P<basis>per annum\b[^\r\n]*)',re.I|re.M)
def amount(s):return format(Decimal(s.replace(',','')).normalize(),'f')
def extract(text):
 matches=list(pattern.finditer(text))
 if len(matches)!=1:return None
 m=matches[0];low=amount(m['low']);high=amount(m['high']) if m['high'] else ''
 if high and Decimal(low)>Decimal(high):raise AssertionError('Invalid source range')
 result='£'+low+('-£'+high if high else '')
 qualifier=('upper limit; ' if m['bound'] and 'up to' in m['bound'].lower() else 'starting amount; ' if m['bound'] else '')
 replacement=m['indent']+'Salary basis: '+qualifier+m['basis']
 return result,text[:m.start()]+replacement+text[m.end():],m[0],replacement
# Narrow regression checks before applying to the supplied file.
assert extract('Salary: £30,000 - £40,000 per annum')[0]=='£30000-£40000'
assert extract('Salary: £24,651.60 – £27,925.20 per annum (Incl. LWA)')[0]=='£24651.6-£27925.2'
assert extract('Salary: up to £95,000 per annum')[0]=='£95000'
for text in ['Salary: £13.43 per hour','Salary: £150 per day','Funding: stipend of £25,555 per annum','Investment of £180 million','Salary: Competitive']:
 assert extract(text) is None
new=[];audit=[]
for rownum,r in enumerate(old,2):
 updated=dict(r);found=extract(r['description'])
 if found:
  assert not r['salaryRange'], 'Do not overwrite an existing salary silently'
  value,description,source,replacement=found
  updated['salaryRange']=value;updated['description']=description
  audit.append({'row':rownum,'jobId':r['jobId'],'company':r['company'],'title':r['title'],'salaryRange':value,'originalSalaryLine':source,'replacementBasisLine':replacement,'note':'Advertised annual salary header used; alternative pro-rata/FTE figures left in description. No hourly conversion or invented endpoint.'})
 new.append(updated)
assert len(audit)==10
csvpath=OUT/'FourCompany-jobs-annual-salary-corrected.csv'
with csvpath.open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(old[0]));w.writeheader();w.writerows(new)
with (OUT/'salary-move-audit.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(audit[0]));w.writeheader();w.writerows(audit)
wb=Workbook();ws=wb.active;ws.title='Corrected jobs';aw=wb.create_sheet('Salary audit');info=wb.create_sheet('Read me')
for sheet,data in [(ws,new),(aw,audit)]:
 sheet.append(list(data[0]))
 for row in data:sheet.append(list(row.values()))
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
 for c in sheet[1]:c.fill=PatternFill('solid',fgColor='17354A');c.font=Font(bold=True,color='FFFFFF')
 for row in sheet.iter_rows(min_row=2):
  sheet.row_dimensions[row[0].row].height=30
  for c in row:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str):c.data_type='s'
 for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
ws.column_dimensions['B'].width=50;ws.column_dimensions['C'].width=90;ws.column_dimensions['D'].width=48
for col in ['F','G','H']:aw.column_dimensions[col].width=70
for a in audit:
 c=ws.cell(a['row'],8);c.fill=PatternFill('solid',fgColor='E1F0E8');c.comment=Comment(a['originalSalaryLine'],'Source salary')
for row in [
 ('Scope','Only annual salary headers moved to salaryRange. All 185 input rows, including 9 unnamed/blank-company rows, are preserved.'),
 ('Results','10 annual salary cells filled: 9 ranges and 1 single upper limit. Remaining 175 salary cells stay blank.'),
 ('Format','£30000-£40000, no commas or extra text. Original GBP retained. Single amounts stay single; no second endpoint invented.'),
 ('Description','Moved monetary amounts removed from Salary headers; Salary basis retains per-annum, upper-limit and London weighting qualifiers. All other description content remains unchanged.'),
 ('Part-time/FTE','Primary advertised salary header used. Alternative pro-rata/FTE salary explanations remain in the original description rather than being merged into an invented range.'),
 ('Excluded','Hourly/daily rates, student stipends, research budgets and funding are not annual employment salaries and were not copied or annualized.'),
 ('Preservation','All other 12 columns and the original CSV are unchanged. Field audit records the exact original and replacement lines.')]:info.append(row)
info.column_dimensions['A'].width=24;info.column_dimensions['B'].width=110
for row in info:
 info.row_dimensions[row[0].row].height=48
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 row[0].font=Font(bold=True)
xlsx=OUT/'FourCompany-jobs-annual-salary-corrected.xlsx';wb.save(xlsx)
# Verify original file, row identity, precisely scoped description edits and both outputs.
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==original_hash
assert len(new)==len(old)==185
byrow={a['row']:a for a in audit}
for i,(before,after) in enumerate(zip(old,new),2):
 assert all(before[k]==after[k] for k in before if k not in ['description','salaryRange'])
 if i not in byrow:assert before==after
 else:
  a=byrow[i];assert after['description']==before['description'].replace(a['originalSalaryLine'],a['replacementBasisLine'],1)
  assert re.fullmatch(r'£\d+(?:\.\d+)?(?:-£\d+(?:\.\d+)?)?',after['salaryRange'])
  assert 'per annum' in a['originalSalaryLine']
assert new[0]['salaryRange']=='' and new[0]['description']==old[0]['description']
assert new[110]['salaryRange']=='' # per-annum PhD stipend is not an employment salary
with csvpath.open(encoding='utf-8-sig',newline='') as f:assert list(csv.DictReader(f))==new
check=load_workbook(xlsx,read_only=True);values=list(check['Corrected jobs'].values)
assert list(values[0])==list(old[0]);assert len(values)==186
for got,want in zip(values[1:],new):assert [v if v is not None else '' for v in got]==list(want.values())
check.close()
summary={'rows':185,'annualSalariesMoved':10,'ranges':9,'singleUpperLimits':1,'salaryBlanksRetained':175,'other12ColumnsUnchanged':True,'unmatched175RowsEntirelyUnchanged':True,'csvWorkbookEquality':True,'originalUnchanged':True,'originalSha256':original_hash,'regressionCasesPassed':8}
(OUT/'validation.json').write_text(json.dumps(summary,indent=2))
print(json.dumps(summary));print([(a['row'],a['salaryRange']) for a in audit])
